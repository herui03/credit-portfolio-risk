"""
verify_claims.py - re-derive every number cited in the challenge docs and assert it.

WHY THIS EXISTS
    Both findings on this project so far were plausible numbers with no error
    attached (challenge-01: a censored denominator; challenge-02: an index read
    against the wrong range). Then the test I wrote to check the SQL files was
    ITSELF broken and returned a false pass - it split statements on ';', hit a
    semicolon inside a header comment, executed a comment-only string, got None
    back, and reported "RUNS OK".

    Three silent-plausible-result failures in a row is a pattern, not bad luck.
    So: every figure quoted in docs/ is asserted here against a live run. If a
    doc and the data disagree, this file fails loudly.

    It also executes the .sql files AS COMMITTED, so an evidence file that does
    not run is a test failure rather than an interview surprise.

Run:  .venv/bin/python python/verify_claims.py
"""

import re
import sys
from pathlib import Path

import duckdb

DB = "data/processed/portfolio.duckdb"
RAW = "data/raw/accepted_2007_to_2018Q4.csv.gz"

results = []


def check(label, actual, expected, tol=0.005):
    """Assert a claimed figure.

    tol is absolute, matching the precision quoted in the docs. Non-numeric types
    fall back to equality - an earlier version coerced everything through float()
    and blew up on a list, which is a test harness failing at the one job it has.
    """
    if isinstance(expected, bool) or isinstance(actual, bool):
        ok = actual == expected
    elif isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        if isinstance(expected, int) and isinstance(actual, int):
            ok = actual == expected
        else:
            ok = abs(float(actual) - float(expected)) <= tol
    else:
        ok = actual == expected
    results.append((ok, label, actual, expected))
    return ok


def run_sql_file(con, path, params=None):
    """Execute a .sql file as committed. Strips comments properly rather than
    naively splitting on ';' - the bug that produced the false pass."""
    text = Path(path).read_text()
    # remove line comments BEFORE splitting, so semicolons inside prose can't split
    stripped = re.sub(r"--[^\n]*", "", text)
    stmts = [s.strip() for s in stripped.split(";") if s.strip()]
    assert len(stmts) == 1, f"{path}: expected 1 statement, found {len(stmts)}"
    return con.execute(stmts[0], params or {}).fetchdf()


def main() -> None:
    con = duckdb.connect(DB, read_only=True)

    # ---------------------------------------------------------------- warehouse
    rows, ids = con.execute("SELECT count(*), count(DISTINCT id) FROM loan").fetchone()
    check("warehouse loan rows", rows, 2_260_668)
    check("warehouse distinct ids (grain)", ids, 2_260_668)

    live_n, live_exp = con.execute(
        "SELECT count(*), sum(out_prncp) FROM loan_live").fetchone()
    check("live book loans", live_n, 907_904)
    check("live book $bn", round(live_exp / 1e9, 3), 9.510, tol=0.001)

    vm = con.execute("SELECT count(DISTINCT vintage_month) FROM loan").fetchone()[0]
    check("distinct vintage months", vm, 139)

    # ------------------------------------------------- challenge-01: the trap
    naive = run_sql_file(con, "sql/naive_vintage_trap.sql").set_index("vintage_yr")

    overall = con.execute("""
        SELECT round(100.0*sum(is_charged_off)/count(*),2)
        FROM loan WHERE loan_status IS NOT NULL
    """).fetchone()[0]
    check("c01 naive default rate overall %", overall, 11.91)

    for yr, exp_naive, exp_curr, exp_resolved in [
        (2014, 17.47, 5.06, 94.68),
        (2015, 18.00, 10.28, 89.18),
        (2016, 15.71, 30.86, 67.47),
        (2017, 8.83, 59.03, 38.17),
        (2018, 1.79, 86.26, 11.37),
    ]:
        check(f"c01 {yr} naive_default_pct", naive.loc[yr, "naive_default_pct"], exp_naive)
        check(f"c01 {yr} still_current_pct", naive.loc[yr, "still_current_pct"], exp_curr)
        check(f"c01 {yr} resolved_pct", naive.loc[yr, "resolved_pct"], exp_resolved)

    for yr, exp in [(2016, 23.28), (2017, 23.12), (2018, 15.75)]:
        check(f"c01 {yr} resolved_only_default_pct",
              naive.loc[yr, "resolved_only_default_pct"], exp)

    # the "8.8x understatement" claim for 2018
    ratio = naive.loc[2018, "resolved_only_default_pct"] / naive.loc[2018, "naive_default_pct"]
    check("c01 2018 understatement multiple", round(ratio, 1), 8.8, tol=0.05)

    # ------------------------------------------- challenge-01: the MOB-matched fix
    vc = run_sql_file(con, "sql/vintage_curve.sql",
                      {"target_mob": 12, "term_filter": "36 months"}).set_index("vintage_yr")

    for yr, exp_n, exp_pct in [
        (2011, 14_101, 4.03), (2012, 43_470, 5.13), (2013, 100_422, 4.23),
        (2014, 162_570, 4.66), (2015, 283_173, 5.27), (2016, 323_495, 6.31),
        (2017, 320_419, 5.66),
    ]:
        check(f"c01 MOB12 {yr} n_loans", int(vc.loc[yr, "n_loans"]), exp_n)
        check(f"c01 MOB12 {yr} cum_default_pct", vc.loc[yr, "cum_default_pct"], exp_pct)

    # THE GATE: 2018 must be absent, not zero
    check("c01 MOB12 2018 absent (gate works)", 2018 not in vc.index, True)

    rel = (vc.loc[2016, "cum_default_pct"] / vc.loc[2011, "cum_default_pct"] - 1) * 100
    check("c01 2011->2016 relative deterioration %", round(rel, 1), 56.6, tol=0.1)

    # ------------------------------------------- challenge-01: MOB reconstruction
    co_total, co_missing = con.execute("""
        SELECT count(*), sum(CASE WHEN last_pymnt_dt IS NULL THEN 1 ELSE 0 END)
        FROM loan WHERE is_charged_off = 1
    """).fetchone()
    check("c01 charged-off loans", co_total, 269_320)
    check("c01 charged-off missing last_pymnt_d", int(co_missing), 2_325)
    check("c01 missing last_pymnt_d %", round(100.0 * co_missing / co_total, 2), 0.86)

    med, mean_, mn, mx, neg = con.execute("""
        SELECT median(mob_stopped_paying), avg(mob_stopped_paying),
               min(mob_stopped_paying), max(mob_stopped_paying),
               sum(CASE WHEN mob_stopped_paying < 0 THEN 1 ELSE 0 END)
        FROM loan WHERE is_charged_off = 1 AND mob_stopped_paying IS NOT NULL
    """).fetchone()
    check("c01 charge-off MOB median", round(med, 1), 14.0)
    check("c01 charge-off MOB mean", round(mean_, 1), 16.1)
    check("c01 charge-off MOB min", int(mn), 0)
    check("c01 charge-off MOB max", int(mx), 66)
    check("c01 charge-off MOB negatives", int(neg), 0)

    # ------------------------------------------- challenge-05: the mob_default rule
    n_null, zero_p, recovered = con.execute("""
        SELECT count(*),
               sum(CASE WHEN total_rec_prncp = 0 THEN 1 ELSE 0 END),
               sum(CASE WHEN mob_default = 0 THEN 1 ELSE 0 END)
        FROM loan WHERE is_charged_off = 1 AND last_pymnt_dt IS NULL
    """).fetchone()
    check("c05 charged-off with no last_pymnt_d", n_null, 2_325)
    check("c05 ALL of them repaid zero principal", int(zero_p), 2_325)
    check("c05 all dated MOB-0 by the rule", int(recovered), 2_325)
    never = con.execute("""
        SELECT count(*) FROM loan
        WHERE is_charged_off=1 AND last_pymnt_dt IS NULL AND total_pymnt = 0
    """).fetchone()[0]
    check("c05 never paid a cent", never, 848)

    # ------------------------------------------- challenge-02: obligor-level HHI
    n_ob, hhi_ob, largest = con.execute("""
        WITH sh AS (SELECT 100.0*out_prncp/sum(out_prncp) OVER () AS pct FROM loan_live)
        SELECT count(*), sum(pct*pct), max(pct) FROM sh
    """).fetchone()
    check("c02 obligors", n_ob, 907_904)
    check("c02 obligor HHI", round(hhi_ob, 4), 0.0179, tol=0.0001)
    check("c02 largest single exposure %", round(largest, 6), 0.000421, tol=0.000001)

    # ------------------------------------------- challenge-02: dimensional HHI
    hh = run_sql_file(con, "sql/concentration_hhi.sql")
    dims = hh.groupby("dimension").first()

    for dim, n, raw, floor, norm in [
        ("addr_state", 50, 507.1, 200.0, 0.031),
        ("purpose", 14, 4074.0, 714.3, 0.362),
        ("grade", 7, 2363.3, 1428.6, 0.109),
        ("term", 2, 5003.0, 5000.0, 0.001),
        ("home_ownership", 5, 4270.0, 2000.0, 0.284),
    ]:
        check(f"c02 {dim} n_buckets", int(dims.loc[dim, "n_buckets"]), n)
        check(f"c02 {dim} hhi_raw", dims.loc[dim, "hhi_raw"], raw, tol=0.05)
        check(f"c02 {dim} hhi_floor", dims.loc[dim, "hhi_floor"], floor, tol=0.05)
        check(f"c02 {dim} hhi_normalised", dims.loc[dim, "hhi_normalised"], norm, tol=0.0005)

    # largest buckets
    st = hh[(hh.dimension == "addr_state")].set_index("rank_by_exposure")
    check("c02 largest state is CA", st.loc[1, "bucket"] == "CA", True)
    check("c02 CA pct_of_book", st.loc[1, "pct_of_book"], 13.27)
    pu = hh[(hh.dimension == "purpose")].set_index("rank_by_exposure")
    check("c02 largest purpose is debt_consolidation",
          pu.loc[1, "bucket"] == "debt_consolidation", True)
    check("c02 debt_consolidation pct_of_book", pu.loc[1, "pct_of_book"], 58.42)
    tm = hh[(hh.dimension == "term")].set_index("rank_by_exposure")
    check("c02 term largest bucket pct", tm.loc[1, "pct_of_book"], 51.23)

    # Top-N curve (the bug that understated Top-10 2x)
    for rk, exp in [(1, 13.27), (3, 29.86), (5, 41.26), (10, 58.10), (20, 79.60)]:
        check(f"c02 top-{rk} states cumulative %", st.loc[rk, "cumulative_pct"], exp)

    # ------------------------------------------- challenge-02: limit framework
    import pandas as pd
    lb = pd.read_csv("outputs/limit_breaches.csv").set_index("limit_id")
    check("c02 limits evaluated", len(lb), 7)
    check("c02 breach count", int((lb.status == "BREACH").sum()), 4)
    for lid, actual, status in [
        ("L-01", 13.275, "BREACH"), ("L-02", 58.422, "BREACH"),
        ("L-03", 58.103, "BREACH"), ("L-04", 6.475, "PASS"),
        ("L-05", 0.362, "BREACH"), ("L-06", 0.031, "PASS"),
        ("L-07", 0.109, "PASS"),
    ]:
        check(f"c02 {lid} actual", lb.loc[lid, "actual"], actual, tol=0.001)
        check(f"c02 {lid} status", lb.loc[lid, "status"] == status, True)

    # ------------------------------------------- challenge-04: the macro sign flip
    import pandas as pd_

    vmac = con.execute("""
        WITH f AS (
            SELECT vintage_yr, mob_observed,
                   CASE WHEN is_charged_off=1 AND mob_default<=12 THEN 1 ELSE 0 END AS d12
            FROM loan WHERE term_clean='36 months' AND loan_status IS NOT NULL
        )
        SELECT vintage_yr, count(*) AS n_loans, 100.0*sum(d12)/count(*) AS mob12
        FROM f WHERE mob_observed>=12 GROUP BY 1 ORDER BY 1
    """).fetchdf()
    mix = con.execute("""
        SELECT vintage_yr,
               100.0*sum(CASE WHEN grade IN ('A','B') THEN 1 ELSE 0 END)/count(*) AS pct_ab,
               100.0*sum(CASE WHEN grade IN ('E','F','G') THEN 1 ELSE 0 END)/count(*) AS pct_efg
        FROM loan WHERE term_clean='36 months' GROUP BY 1
    """).fetchdf()
    un = pd_.read_csv("data/seed/fred_unrate.csv", parse_dates=["observation_date"])
    un["vintage_yr"] = un["observation_date"].dt.year
    un = un.groupby("vintage_yr", as_index=False)["UNRATE"].mean().round(2)
    mac = vmac.merge(mix, on="vintage_yr").merge(un, on="vintage_yr")

    check("c04 macro observations available", len(mac), 11)
    check("c04 corr(default, unemployment)", round(mac.mob12.corr(mac.UNRATE), 3), -0.521, tol=0.001)
    check("c04 corr(default, pct grade A/B)", round(mac.mob12.corr(mac.pct_ab), 3), -0.908, tol=0.001)
    check("c04 volume growth 2007->2017 (x)",
          round(mac.n_loans.iloc[-1] / mac.n_loans.iloc[0]), 531, tol=1)
    m = mac.set_index("vintage_yr")
    for yr, unrate, ab, efg in [(2007, 4.62, 29.2, 31.0), (2011, 8.93, 73.1, 2.4)]:
        check(f"c04 {yr} unemployment", m.loc[yr, "UNRATE"], unrate)
        check(f"c04 {yr} pct grade A/B", round(m.loc[yr, "pct_ab"], 1), ab, tol=0.05)
        check(f"c04 {yr} pct grade E/F/G", round(m.loc[yr, "pct_efg"], 1), efg, tol=0.05)

    # rating drift: every grade deteriorated 2013 -> 2016
    drift = con.execute("""
        WITH f AS (
            SELECT vintage_yr, grade, mob_observed,
                   CASE WHEN is_charged_off=1 AND mob_default<=12 THEN 1 ELSE 0 END AS d12
            FROM loan WHERE term_clean='36 months' AND loan_status IS NOT NULL AND grade IS NOT NULL
        )
        SELECT vintage_yr, grade, round(100.0*sum(d12)/count(*),2) AS pct
        FROM f WHERE mob_observed>=12 AND vintage_yr IN (2013, 2016)
        GROUP BY 1,2
    """).fetchdf().pivot(index="grade", columns="vintage_yr", values="pct")
    for g, v13, v16 in [("A", 1.16, 1.75), ("B", 2.88, 4.15), ("C", 5.25, 7.90), ("D", 7.96, 12.62)]:
        check(f"c04 grade {g} 2013 MOB-12", drift.loc[g, 2013], v13)
        check(f"c04 grade {g} 2016 MOB-12", drift.loc[g, 2016], v16)
        check(f"c04 grade {g} drift is deterioration", bool(drift.loc[g, 2016] > drift.loc[g, 2013]), True)

    # LGD basis - the 64% vs 89% error
    lgd_f, lgd_e = con.execute("""
        SELECT 100.0*(sum(funded_amnt)-sum(total_rec_prncp)-sum(recoveries))/sum(funded_amnt),
               100.0*(sum(funded_amnt)-sum(total_rec_prncp)-sum(recoveries))
                   /(sum(funded_amnt)-sum(total_rec_prncp))
        FROM loan WHERE is_charged_off=1
    """).fetchone()
    check("c04 LGD vs funded (wrong basis)", round(lgd_f, 2), 64.02)
    check("c04 LGD vs EAD (correct basis)", round(lgd_e, 2), 89.17)

    # live-book seasoning - the declared conservative bias
    med_mob, mean_mob, pct_past = con.execute("""
        SELECT median(mob_observed), avg(mob_observed),
               100.0*sum(CASE WHEN mob_observed>=12 THEN 1 ELSE 0 END)/count(*)
        FROM loan_live
    """).fetchone()
    check("c04 live book median MOB", round(med_mob, 1), 12.0)
    check("c04 live book mean MOB", round(mean_mob, 1), 14.6)
    check("c04 live book % past MOB 12", round(pct_past, 1), 51.8)

    # stress results
    sr = pd_.read_csv("outputs/stress_results.csv").set_index("scenario_id")
    for sid, mult, rate in [("S-00", 1.00, 5.76), ("S-01", 1.50, 8.65), ("S-02", 2.33, 13.43)]:
        check(f"c04 {sid} pd_multiplier", sr.loc[sid, "pd_multiplier"], mult)
        check(f"c04 {sid} loss_rate_pct", sr.loc[sid, "loss_rate_pct"], rate)
    # reverse stress must invert the Severe scenario exactly
    base_el = sr.loc["S-00", "expected_loss_mm"]
    check("c04 reverse stress inverts S-02 (13.43% -> 2.33x)",
          round(13.43 / sr.loc["S-00", "loss_rate_pct"], 2), 2.33, tol=0.01)

    # ------------------------------------------- EWI module (challenge-06)
    fpd = run_sql_file(con, "sql/ewi_fpd.sql").set_index("vintage_yr")
    for yr, n_loans, n_fpd, pct in [
        (2013, 134_814, 16, 0.012), (2014, 235_629, 36, 0.015),
        (2015, 421_095, 68, 0.016), (2016, 434_407, 140, 0.032),
        (2017, 443_579, 210, 0.047), (2018, 238_636, 175, 0.073),
    ]:
        check(f"c06 FPD {yr} n_loans", int(fpd.loc[yr, "n_loans"]), n_loans)
        check(f"c06 FPD {yr} n_fpd", int(fpd.loc[yr, "n_fpd"]), n_fpd)
        check(f"c06 FPD {yr} fpd_pct", fpd.loc[yr, "fpd_pct"], pct, tol=0.0005)

    # never-paid population is cleanly identifiable
    never, never_co, never_open, has_date = con.execute("""
        SELECT sum(CASE WHEN total_pymnt=0 THEN 1 ELSE 0 END),
               sum(CASE WHEN total_pymnt=0 AND is_charged_off=1 THEN 1 ELSE 0 END),
               sum(CASE WHEN total_pymnt=0 AND is_charged_off=0 THEN 1 ELSE 0 END),
               sum(CASE WHEN total_pymnt=0 AND last_pymnt_dt IS NOT NULL THEN 1 ELSE 0 END)
        FROM loan
    """).fetchone()
    check("c06 never-paid loans", int(never), 949)
    check("c06 never-paid AND charged off", int(never_co), 848)
    check("c06 never-paid not yet labelled", int(never_open), 101)
    check("c06 never-paid with a payment date (must be 0)", int(has_date), 0)

    # the gate is protective, not material - assert the honest claim
    ungated = con.execute("""
        SELECT round(100.0*sum(CASE WHEN total_pymnt=0 AND is_charged_off=1 THEN 1 ELSE 0 END)
               /count(*),3) FROM loan WHERE vintage_yr=2018
    """).fetchone()[0]
    check("c06 2018 FPD ungated", ungated, 0.070)
    check("c06 2018 FPD gated (gate moves it +0.003pp only)", fpd.loc[2018, "fpd_pct"], 0.073)

    # rating drift closes the L-04/L-07 gap
    dr = run_sql_file(con, "sql/ewi_rating_drift.sql")
    d16 = dr[dr.vintage_yr == 2016].set_index("grade")
    for g, base, val, drift in [
        ("A", 1.27, 1.75, 36.9), ("B", 3.12, 4.15, 33.0),
        ("C", 5.99, 7.90, 31.8), ("D", 9.22, 12.62, 36.8),
    ]:
        check(f"c06 drift {g} baseline", d16.loc[g, "baseline_pct"], base)
        check(f"c06 drift {g} 2016", d16.loc[g, "mob12_pct"], val)
        check(f"c06 drift {g} pct", d16.loc[g, "drift_pct"], drift, tol=0.05)

    # THE HEADLINE: the slow panel is silent on the largest vintage
    panel = pd_.read_csv("outputs/ewi_panel.csv")
    e03_2018 = panel[(panel.ewi_id == "E-03") & (panel.vintage_yr == 2018)].iloc[0]
    check("c06 E-03 has NO SIGNAL for 2018", "NO SIGNAL" in e03_2018.status, True)
    e01_2018 = panel[(panel.ewi_id == "E-01") & (panel.vintage_yr == 2018)].iloc[0]
    check("c06 E-01 is RED for 2018", e01_2018.status == "RED", True)
    check("c06 E-01 2018 ratio to baseline", e01_2018.ratio_or_drift, 5.093, tol=0.001)
    n2018 = con.execute("SELECT count(*) FROM loan WHERE vintage_yr=2018").fetchone()[0]
    check("c06 2018 vintage size", n2018, 495_242)

    # ------------------------------------------- Excel pack (challenge-07)
    import openpyxl
    XL = "outputs/risk_committee_pack.xlsx"
    check("c07 pack exists", Path(XL).exists(), True)
    wbk = openpyxl.load_workbook(XL)
    check("c07 pack sheets", wbk.sheetnames ==
          ["Cover", "Limits", "Concentration", "EWI Panel", "Stress"], True)

    # SOURCE values must match the verified CSVs cell-by-cell. The workbook is a
    # rendering; the CSVs are the source of truth. If they disagree the pack is wrong.
    lim_csv = pd_.read_csv("outputs/limit_breaches.csv")
    wsl = wbk["Limits"]
    mismatches = 0
    for i, row in lim_csv.iterrows():
        r = i + 2
        if abs(wsl.cell(r, 5).value - row.actual) > 1e-9: mismatches += 1
        if abs(wsl.cell(r, 6).value - row.threshold) > 1e-9: mismatches += 1
    check("c07 Limits source cells match CSV", mismatches, 0)

    st_csv = pd_.read_csv("outputs/stress_results.csv")
    wss = wbk["Stress"]
    sm = 0
    for i, row in st_csv.iterrows():
        r = i + 2
        if abs(wss.cell(r, 5).value - row.loss_rate_pct) > 1e-9: sm += 1
        if abs(wss.cell(r, 3).value - row.pd_multiplier) > 1e-9: sm += 1
    check("c07 Stress source cells match CSV", sm, 0)

    # DERIVED cells are formulas. The harness can assert the TEXT and never the
    # RESULT - no calculation engine runs here. That boundary is the finding.
    check("c07 utilisation is a formula, not a value", wsl.cell(2, 7).value, "=E2/F2*100")
    check("c07 status is a formula, not a value",
          wsl.cell(2, 8).value, '=IF(E2>F2,"BREACH","PASS")')
    # prove the claim: a formula's cached value is unreadable from this pipeline
    wbk_v = openpyxl.load_workbook(XL, data_only=True)
    check("c07 formula RESULT is not verifiable (reads 0/None)",
          wbk_v["Limits"].cell(2, 7).value in (0, None), True)

    # NO SIGNAL must render as text, never as a blank or a zero
    wse = wbk["EWI Panel"]
    nosig_rows = [r for r in range(2, 12) if "NO SIGNAL" in str(wse.cell(r, 8).value)]
    check("c07 exactly one NO SIGNAL row", len(nosig_rows), 1)
    check("c07 NO SIGNAL value cell is a dash, not 0",
          wse.cell(nosig_rows[0], 5).value, "\u2014")

    # the pack must carry its own caveats
    cover_text = " ".join(str(c.value) for row in wbk["Cover"].iter_rows() for c in row)
    for phrase in ["ILLUSTRATIVE", "NO SIGNAL' IS NOT 'GREEN", "OVERSTATES LOSS",
                   "NO MACRO MODEL IS FITTED", "NOT VERIFIED BY THE PIPELINE"]:
        check(f"c07 Cover carries caveat: {phrase[:28]}", phrase in cover_text, True)

    # ------------------------------------------- Tableau extracts (challenge-08)
    # The BUILD_GUIDE is untestable. The extracts are not, so they get asserted.
    grains = {
        "tableau/extract_vintage.csv": 36, "tableau/extract_concentration.csv": 78,
        "tableau/extract_ewi.csv": 10, "tableau/extract_limits.csv": 7,
        "tableau/extract_stress.csv": 3,
    }
    for f_, n_ in grains.items():
        check(f"c08 {Path(f_).name} rows", len(pd_.read_csv(f_)), n_)

    # extracts must not drift from the verified outputs they render
    for src, ext, key in [
        ("outputs/concentration.csv", "tableau/extract_concentration.csv", "exposure_mm"),
        ("outputs/limit_breaches.csv", "tableau/extract_limits.csv", "actual"),
        ("outputs/stress_results.csv", "tableau/extract_stress.csv", "loss_rate_pct"),
    ]:
        a = pd_.read_csv(src)[key].round(6).tolist()
        b = pd_.read_csv(ext)[key].round(6).tolist()
        check(f"c08 {Path(ext).name} matches its source", a == b, True)

    # THE FAN TRAP: assert the hazard the guide warns about is real
    lim_ = pd_.read_csv("outputs/limit_breaches.csv")
    con_ = pd_.read_csv("outputs/concentration.csv")
    fan = lim_.merge(con_, on="dimension", how="inner")
    check("c08 naive join fans 7x78 into", len(fan), 192)
    ca_true = con_[(con_.dimension == "addr_state") & (con_.bucket == "CA")].exposure_mm.iloc[0]
    ca_fan = fan[(fan.dimension == "addr_state") & (fan.bucket == "CA")].exposure_mm.sum()
    check("c08 true CA exposure $mm", float(ca_true), 1262.5, tol=0.05)
    check("c08 CA after fan trap $mm", float(ca_fan), 3787.5, tol=0.05)
    check("c08 fan trap overstates 3x", round(ca_fan / ca_true), 3)

    # THE NULL IS LOAD-BEARING: 2018 MOB-12 must stay NULL, never 0
    vx = pd_.read_csv("tableau/extract_vintage.csv")
    m18 = vx[(vx.vintage_yr == 2018) & (vx.metric == "mob12_default_pct")].iloc[0]
    check("c08 2018 MOB-12 is NULL in the extract", bool(pd_.isna(m18.value)), True)
    check("c08 the NULL carries its own explanation",
          "NOT OBSERVED" in str(m18.observability_status), True)
    check("c08 extract_vintage has exactly one NULL", int(vx.value.isna().sum()), 1)

    dd = pd_.read_csv("tableau/DATA_DICTIONARY.csv")
    check("c08 data dictionary fields documented", len(dd), 42)

    # ------------------------------------------- README (it is read first, so it is asserted first)
    readme = Path("README.md").read_text()
    # every figure the README quotes must appear verbatim and be true elsewhere in this file
    for claim in [
        "2,260,668", "907,904", "$9.510bn", "2,260,701", "56.6%",
        "18.00", "1.79", "86.26", "4.03", "6.31", "5.27", "4.23", "5.66",
        "58.42%", "58.10%", "0.073%", "5.093x", "31.8% to +36.9%",
        "5.76% | Adverse 8.65% | Severe 13.43%",
        "\u22120.521", "\u22120.908", "2,325", "495,242", "3,787.5mm",
        "13.275", "99.999", "192 rows instead of 78",     ]:
        check(f"readme quotes {claim[:34]}", claim in readme, True)

    # the README must NOT claim the dashboard exists
    check("readme does not claim Tableau is built",
          "the dashboard is not built yet" in readme.lower()
          or "The dashboard itself does not exist yet" in readme, True)
    # the README must carry the illustrative caveat
    check("readme carries ILLUSTRATIVE caveat", "ILLUSTRATIVE" in readme, True)
    # a hardcoded assertion count in the README would be stale the next time this
    # file grows. It must not come back.
    import re as _re2
    check("readme hardcodes no assertion count",
          _re2.search(r"\d+\s*assertions|\d+/\d+ claims verified", readme) is None, True)
    # counts quoted in the README must be real
    import glob as _g
    # NOT line counts, and NOT an assertion count.
    # An earlier README quoted "1,716 python lines" and the act of adding that very
    # assertion pushed it to 1,743. It also quoted "242 assertions" while the real
    # number had moved to 277 - and the check only verified the STRING was present,
    # not that it was TRUE, which is a test that tests nothing (challenge-03).
    # Any count of your own artifacts, quoted in your own artifacts, checked by a
    # counter that is itself an artifact, rots on every edit. So the README no longer
    # hardcodes either; `make verify` prints the live number instead.
    # File counts are stable facts about the deliverable and are safe to assert.
    check("readme module count", len(_g.glob("python/*.py")), 8)
    check("readme sql file count", len(_g.glob("sql/*.sql")), 5)
    check("readme challenge count", len(_g.glob("docs/challenge-*.md")), 8)

    # ------------------------------------------- evidence files must exist
    for f in [
        "sql/vintage_curve.sql", "sql/naive_vintage_trap.sql",
        "sql/concentration_hhi.sql", "config/risk_limits.csv",
        "python/build_warehouse.py", "python/limit_monitor.py",
        "python/profile_lendingclub.py", "outputs/limit_breaches.csv",
        "outputs/vintage_naive_vs_matched.csv",
        "docs/assumptions_and_limitations.md",
        "docs/challenge-01-maturity-bias.md", "docs/challenge-02-hhi-cardinality.md",
        "docs/challenge-03-the-test-that-tested-nothing.md",
        "docs/challenge-04-the-recession-that-helped.md",
        "python/stress_test.py", "config/stress_scenarios.csv",
        "outputs/stress_results.csv", "data/seed/fred_unrate.csv",
        "docs/challenge-05-unknown-is-not-zero-again.md",
        "docs/challenge-06-the-silent-panel.md",
        "sql/ewi_fpd.sql", "sql/ewi_rating_drift.sql",
        "python/ewi_monitor.py", "config/ewi_thresholds.csv", "outputs/ewi_panel.csv",
        "docs/challenge-07-the-logic-that-left-the-harness.md",
        "python/build_excel_pack.py", "outputs/risk_committee_pack.xlsx",
        "outputs/concentration.csv",
        "docs/challenge-08-the-fan-trap.md",
        "python/build_tableau_extracts.py", "tableau/BUILD_GUIDE.md",
        "tableau/DATA_DICTIONARY.csv", "tableau/extract_vintage.csv",
        "README.md", "requirements.txt", "Makefile",
    ]:
        check(f"evidence file exists: {f}", Path(f).exists(), True)

    # ------------------------------------------------------------------ report
    failed = [r for r in results if not r[0]]
    for ok, label, actual, expected in results:
        if not ok:
            print(f"  FAIL  {label}: got {actual!r}, doc claims {expected!r}")
    print(f"\n{len(results) - len(failed)}/{len(results)} claims verified")
    if failed:
        print(f"{len(failed)} FAILED - a doc and the data disagree. Fix one of them.")
        sys.exit(1)
    print("All figures in docs/ re-derived from a live run against committed code.")
    con.close()


if __name__ == "__main__":
    main()
